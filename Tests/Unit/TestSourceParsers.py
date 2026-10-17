from io import BytesIO

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfWriter

from NoteAgent.BusinessModules.SourceFiles.SourceFileParsers import parse_source, ParseError


def test_text_has_locator_and_no_truncation():
    parts = parse_source("text/plain", b"one\n\ntwo")
    assert [part.text for part in parts] == ["one", "two"]
    assert parts[0].locator == "paragraph 1"


def test_docx_preserves_paragraph_and_table_order():
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "table"
    doc.add_paragraph("after")
    stream = BytesIO()
    doc.save(stream)
    parts = parse_source("application/vnd.openxmlformats-officedocument.wordprocessingml.document", stream.getvalue())
    assert [part.text for part in parts] == ["before", "table", "after"]


def test_image_requires_vision_and_enforces_pixel_limit():
    stream = BytesIO()
    Image.new("RGB", (2, 2)).save(stream, format="PNG")
    with pytest.raises(ParseError):
        parse_source("image/png", stream.getvalue())
    parts = parse_source("image/png", stream.getvalue(), vision=lambda data, media: "diagram")
    assert parts[0].text == "diagram"
    with pytest.raises(ParseError):
        parse_source("image/png", stream.getvalue(), vision=lambda data, media: "x", max_pixels=3)


def test_blank_pdf_page_uses_vision_with_page_locator():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(ParseError):
        parse_source("application/pdf", stream.getvalue())
    parts = parse_source("application/pdf", stream.getvalue(), vision=lambda data, media: "a diagram")
    assert parts[0].locator == "page 1"
    assert parts[0].text == "a diagram"
